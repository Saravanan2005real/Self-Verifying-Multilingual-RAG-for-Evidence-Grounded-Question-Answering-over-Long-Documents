import sys
import json
import re
from pathlib import Path
from unittest.mock import patch, MagicMock
import pytest
import docx
import fitz
import requests

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

from src.document_loader import DocumentPage, DocumentLoader
from src.chunking import DocumentChunk, MeaningfulSemanticChunker
from src.semantic_chunker import OllamaSemanticAdvisor, LlamaSemanticChunker
from main import process_document

SAMPLE_DIR = BASE_DIR / "tests" / "sample_docs"
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)


# =====================================================================
# TEST 1: Llama Semantic Chunking (Decision Handling)
# =====================================================================
def test_01_llama_semantic_chunking_decisions():
    """Verify that a valid BREAK decision from Llama starts a new chunk when min_words is satisfied."""
    advisor = OllamaSemanticAdvisor(base_url="http://mock-ollama:11434")
    chunker = LlamaSemanticChunker(advisor=advisor, min_words=50, target_words=100, max_words=200, enabled=True)

    p1 = "Paragraph 1 introduces the fundamental principles of remote sensing for surface hydrology. " * 5
    p2 = "Paragraph 2 shifts topic completely to machine learning models for classification. " * 5

    pages = [DocumentPage(page_number=1, content=f"{p1}\n\n{p2}", source="test_doc", source_type="docx")]

    # Mock Ollama returning a BREAK for paragraph 1
    mock_response = MagicMock()
    mock_response.status_code = 200
    mock_response.json.return_value = {
        "message": {
            "content": json.dumps({
                "decisions": [
                    {"paragraph_index": 0, "decision": "CONTINUE"},
                    {"paragraph_index": 1, "decision": "BREAK"}
                ]
            })
        }
    }

    with patch("requests.get") as mock_get, patch("requests.post") as mock_post:
        mock_get.return_value.status_code = 200
        mock_get.return_value.json.return_value = {"models": [{"name": "llama3.2:latest"}]}
        mock_post.return_value = mock_response

        chunks = chunker.chunk_pages(pages)

    assert len(chunks) == 2, f"Expected 2 chunks due to BREAK decision, got {len(chunks)}"
    assert "Paragraph 1" in chunks[0].text
    assert "Paragraph 2" in chunks[1].text
    assert chunks[0].metadata.get("chunking_strategy") == "llama3.2_semantic"


# =====================================================================
# TEST 2: Ollama Success
# =====================================================================
def test_02_ollama_success():
    """Verify successful Ollama communication and structured decision parsing."""
    advisor = OllamaSemanticAdvisor(base_url="http://mock-ollama:11434")
    
    mock_resp = MagicMock()
    mock_resp.status_code = 200
    mock_resp.json.return_value = {
        "message": {
            "content": json.dumps({
                "decisions": [
                    {"paragraph_index": 1, "decision": "CONTINUE"},
                    {"paragraph_index": 2, "decision": "BREAK"}
                ]
            })
        }
    }

    with patch("requests.post", return_value=mock_resp):
        paragraphs = [
            {"text": "Sample para 0", "word_count": 3},
            {"text": "Sample para 1", "word_count": 3},
            {"text": "Sample para 2", "word_count": 3},
        ]
        decisions = advisor.evaluate_boundaries_window(paragraphs, start_index=0)

    assert decisions == {1: "CONTINUE", 2: "BREAK"}


# =====================================================================
# TEST 3: Ollama Unavailable (ConnectionError)
# =====================================================================
def test_03_ollama_unavailable():
    """Verify that when Ollama is offline (ConnectionError), the system falls back without crashing."""
    advisor = OllamaSemanticAdvisor(base_url="http://offline-ollama:11434")
    chunker = LlamaSemanticChunker(advisor=advisor, min_words=30, target_words=60, max_words=120, enabled=True)

    pages = [DocumentPage(page_number=1, content="This is page text that should be chunked deterministically.", source="doc", source_type="docx")]

    with patch("requests.get", side_effect=requests.exceptions.ConnectionError("Connection refused")):
        chunks = chunker.chunk_pages(pages)

    assert len(chunks) >= 1
    assert chunks[0].metadata.get("chunking_strategy") == "deterministic_fallback"


# =====================================================================
# TEST 4: Ollama Timeout
# =====================================================================
def test_04_ollama_timeout():
    """Verify that when Ollama times out, the system triggers graceful fallback."""
    advisor = OllamaSemanticAdvisor(base_url="http://slow-ollama:11434")
    chunker = LlamaSemanticChunker(advisor=advisor, min_words=30, target_words=60, max_words=120, enabled=True)

    p1 = "This is paragraph one with enough content describing the initial conditions."
    p2 = "This is paragraph two with follow-up content detailing the observed results."
    pages = [DocumentPage(page_number=1, content=f"{p1}\n\n{p2}", source="doc", source_type="docx")]

    with patch("requests.get") as mock_get, patch("requests.post", side_effect=requests.exceptions.Timeout("Request timed out")):
        mock_get.return_value.status_code = 200
        mock_get.return_value.json.return_value = {"models": [{"name": "llama3.2:latest"}]}
        chunks = chunker.chunk_pages(pages)

    assert len(chunks) >= 1
    assert chunks[0].metadata.get("chunking_strategy") == "deterministic_fallback"


# =====================================================================
# TEST 5: Invalid Llama Response (Malformed JSON)
# =====================================================================
def test_05_invalid_llama_response():
    """Verify that invalid/malformed JSON from Llama does not crash the system and triggers fallback."""
    advisor = OllamaSemanticAdvisor(base_url="http://mock-ollama:11434")
    chunker = LlamaSemanticChunker(advisor=advisor, min_words=30, target_words=60, max_words=120, enabled=True)

    p1 = "Some document text in paragraph one that will be evaluated for transitions."
    p2 = "Some follow-up text in paragraph two that will encounter a malformed LLM response."
    pages = [DocumentPage(page_number=1, content=f"{p1}\n\n{p2}", source="doc", source_type="docx")]

    mock_resp = MagicMock()
    mock_resp.status_code = 200
    mock_resp.json.return_value = {
        "message": {"content": "INVALID_JSON_HERE_NOT_A_DICT"}
    }

    with patch("requests.get") as mock_get, patch("requests.post", return_value=mock_resp):
        mock_get.return_value.status_code = 200
        mock_get.return_value.json.return_value = {"models": [{"name": "llama3.2:latest"}]}
        chunks = chunker.chunk_pages(pages)

    assert len(chunks) >= 1
    assert chunks[0].metadata.get("chunking_strategy") == "deterministic_fallback"


# =====================================================================
# TEST 6: Deterministic Fallback Strategy
# =====================================================================
def test_06_deterministic_fallback():
    """Verify that deterministic fallback preserves 100% of text and tags metadata correctly."""
    chunker = LlamaSemanticChunker(enabled=False, min_words=20, target_words=40, max_words=80)
    p_text = "Deterministic chunking always delivers reproducible results regardless of LLM availability."
    pages = [DocumentPage(page_number=1, content=p_text, source="doc", source_type="docx")]

    chunks = chunker.chunk_pages(pages)
    assert len(chunks) >= 1
    assert p_text in chunks[0].text
    assert chunks[0].metadata.get("chunking_strategy") == "deterministic_disabled"


# =====================================================================
# TEST 7: First Paragraph Preservation
# =====================================================================
def test_07_first_paragraph_preservation():
    """Verify that the very first paragraph of the document is preserved in chunk 0."""
    chunker = LlamaSemanticChunker(enabled=True, min_words=30, target_words=60, max_words=120)
    first_p = "CRITICAL_FIRST_PARAGRAPH: This study begins with an in-depth exploration of hydrological systems."
    second_p = "The second paragraph transitions into satellite multispectral analysis methodology."
    pages = [DocumentPage(page_number=1, content=f"{first_p}\n\n{second_p}", source="doc", source_type="docx")]

    with patch.object(chunker.advisor, "is_available", return_value=False):
        chunks = chunker.chunk_pages(pages)

    assert len(chunks) >= 1
    assert first_p in chunks[0].text


# =====================================================================
# TEST 8: Last Paragraph Preservation
# =====================================================================
def test_08_last_paragraph_preservation():
    """Verify that the final paragraph of the document is completely flushed into the chunks."""
    chunker = LlamaSemanticChunker(enabled=True, min_words=30, target_words=60, max_words=120)
    paras = [f"Intermediate paragraph {i} explaining experimental parameters in detail." for i in range(1, 10)]
    final_p = "FINAL_CONCLUSION_PARAGRAPH: All hypotheses were validated with 99.8% statistical confidence."
    full_text = "\n\n".join(paras + [final_p])
    pages = [DocumentPage(page_number=1, content=full_text, source="doc", source_type="docx")]

    with patch.object(chunker.advisor, "is_available", return_value=False):
        chunks = chunker.chunk_pages(pages)

    all_text = " ".join(c.text for c in chunks)
    assert final_p in all_text


# =====================================================================
# TEST 9: Paragraph Ordering
# =====================================================================
def test_09_paragraph_ordering():
    """Verify that paragraphs appear in strict ascending sequential order across all chunks."""
    chunker = LlamaSemanticChunker(enabled=True, min_words=40, target_words=80, max_words=160)
    paras = [f"SEQ_TAG_{i:02d}: Document section content description {i}." * 8 for i in range(1, 8)]
    pages = [DocumentPage(page_number=1, content="\n\n".join(paras), source="doc", source_type="docx")]

    with patch.object(chunker.advisor, "is_available", return_value=False):
        chunks = chunker.chunk_pages(pages)

    last_found = -1
    for chunk in chunks:
        for i in range(1, 8):
            tag = f"SEQ_TAG_{i:02d}"
            if tag in chunk.text:
                assert i >= last_found, f"Ordering violation: {tag} appeared out of sequence"
                last_found = i


# =====================================================================
# TEST 10: No Duplicate Paragraphs
# =====================================================================
def test_10_no_duplicate_paragraphs():
    """Verify that non-overlapping paragraph contents are not duplicated across chunks."""
    chunker = LlamaSemanticChunker(enabled=True, min_words=30, target_words=60, max_words=120, overlap_words=0)
    paras = [f"UNIQUE_PARAGRAPH_{i}: Distinct topic explanation number {i}." for i in range(1, 6)]
    pages = [DocumentPage(page_number=1, content="\n\n".join(paras), source="doc", source_type="docx")]

    with patch.object(chunker.advisor, "is_available", return_value=False):
        chunks = chunker.chunk_pages(pages)

    for i in range(1, 6):
        tag = f"UNIQUE_PARAGRAPH_{i}"
        occurrences = sum(1 for c in chunks if tag in c.text)
        assert occurrences == 1, f"Paragraph {tag} was duplicated {occurrences} times"


# =====================================================================
# TEST 11: No Data Loss (100% Text Coverage)
# =====================================================================
def test_11_no_data_loss():
    """Verify that every single input paragraph exists in the output chunks."""
    chunker = LlamaSemanticChunker(enabled=True, min_words=25, target_words=50, max_words=100)
    paras = [f"Paragraph index {i} detailing analytical evaluation protocol." for i in range(1, 8)]
    pages = [DocumentPage(page_number=1, content="\n\n".join(paras), source="doc", source_type="docx")]

    with patch.object(chunker.advisor, "is_available", return_value=False):
        chunks = chunker.chunk_pages(pages)

    combined_text = " ".join(c.text for c in chunks)
    for p in paras:
        assert p in combined_text, f"Data lost: '{p}' not found in any chunk"


# =====================================================================
# TEST 12: Short Paragraphs Preserved
# =====================================================================
def test_12_short_paragraphs_preserved():
    """Verify that short titles, metadata lines, and bullet items are never silently dropped."""
    chunker = LlamaSemanticChunker(enabled=True, min_words=30, target_words=60, max_words=120)
    short_title = "Date: 13 April 2026"
    bullet_item = "• Item A"
    body_text = "This is a detailed paragraph explaining the contextual background of the project in depth."
    pages = [DocumentPage(page_number=1, content=f"{short_title}\n\n{bullet_item}\n\n{body_text}", source="doc", source_type="docx")]

    with patch.object(chunker.advisor, "is_available", return_value=False):
        chunks = chunker.chunk_pages(pages)

    combined_text = " ".join(c.text for c in chunks)
    assert short_title in combined_text
    assert bullet_item in combined_text


# =====================================================================
# TEST 13: Long Paragraphs Split Safely
# =====================================================================
def test_13_long_paragraphs_split_safely():
    """Verify that a paragraph exceeding max_words is split safely without losing a single word."""
    chunker = LlamaSemanticChunker(enabled=True, min_words=30, target_words=50, max_words=80)
    sentences = [f"Sentence {i} gives crucial scientific evidence regarding climate variance." for i in range(1, 20)]
    long_para = " ".join(sentences)
    pages = [DocumentPage(page_number=1, content=long_para, source="doc", source_type="docx")]

    with patch.object(chunker.advisor, "is_available", return_value=False):
        chunks = chunker.chunk_pages(pages)

    assert len(chunks) >= 2, "Long paragraph should have been split into multiple chunks"
    combined_text = " ".join(c.text for c in chunks)
    for s in sentences:
        assert s in combined_text, f"Sentence lost during split: {s}"


# =====================================================================
# TEST 14: English Text Preservation
# =====================================================================
def test_14_english_text_preservation():
    """Verify English document with headings and technical prose."""
    chunker = LlamaSemanticChunker(enabled=True)
    text = (
        "1. Introduction and Objectives\n\n"
        "This research paper evaluates digital image processing techniques for geospatial analytics.\n\n"
        "The subsequent sections discuss convolutional filtering and threshold segmentation."
    )
    pages = [DocumentPage(page_number=1, content=text, source="doc", source_type="docx")]

    with patch.object(chunker.advisor, "is_available", return_value=False):
        chunks = chunker.chunk_pages(pages)

    assert len(chunks) >= 1
    assert "1. Introduction and Objectives" in chunks[0].text
    assert "This research paper evaluates" in chunks[0].text
    assert chunks[0].language == "en"


# =====================================================================
# TEST 15: Hindi Unicode Preservation
# =====================================================================
def test_15_hindi_unicode_preservation():
    """Verify Hindi Devanagari script and purna viram (।) sentence terminators."""
    chunker = LlamaSemanticChunker(enabled=True)
    hindi_text = (
        "१. परिचय एवं उद्देश्य\n\n"
        "यह शोध जल निकायों के विश्लेषण पर केंद्रित है। तिरुपति में उपग्रह चित्रों का अध्ययन किया गया है।\n\n"
        "गूगल अर्थ इंजन के माध्यम से डेटा संकलित किया गया है।"
    )
    pages = [DocumentPage(page_number=1, content=hindi_text, source="doc", source_type="docx")]

    with patch.object(chunker.advisor, "is_available", return_value=False):
        chunks = chunker.chunk_pages(pages)

    assert len(chunks) >= 1
    assert "परिचय एवं उद्देश्य" in chunks[0].text
    assert "जल निकायों के विश्लेषण" in chunks[0].text
    assert chunks[0].language == "hi"


# =====================================================================
# TEST 16: Tamil Unicode Preservation
# =====================================================================
def test_16_tamil_unicode_preservation():
    """Verify Tamil Unicode script preservation without character dropping or corruption."""
    chunker = LlamaSemanticChunker(enabled=True)
    tamil_text = (
        "1. அறிமுகம் மற்றும் நோக்கங்கள்\n\n"
        "இந்த ஆய்வு செயற்கைக்கோள் படங்களைப் பயன்படுத்தி நீர்நிலைகளின் மாற்றங்களை ஆராய்கிறது.\n\n"
        "திருப்பதியில் உள்ள முக்கியமான ஏரிகள் மற்றும் குளங்கள் பகுப்பாய்வு செய்யப்பட்டன."
    )
    pages = [DocumentPage(page_number=1, content=tamil_text, source="doc", source_type="docx")]

    with patch.object(chunker.advisor, "is_available", return_value=False):
        chunks = chunker.chunk_pages(pages)

    assert len(chunks) >= 1
    assert "அறிமுகம் மற்றும் நோக்கங்கள்" in chunks[0].text
    assert "செயற்கைக்கோள் படங்களைப்" in chunks[0].text
    assert chunks[0].language == "ta"


# =====================================================================
# TEST 17: DOCX Processing with Semantic Chunker
# =====================================================================
def test_17_docx_processing():
    """Verify end-to-end process_document on a DOCX file with semantic chunking."""
    doc_path = SAMPLE_DIR / "test_17_sem.docx"
    doc = docx.Document()
    doc.add_heading("Section 1: Data Acquisition", level=1)
    doc.add_paragraph("Remote sensing data is downloaded from Copernicus open access hub.")
    doc.add_paragraph("The bands are composited into false color infrared.")
    doc.save(str(doc_path))

    with patch("src.semantic_chunker.OllamaSemanticAdvisor.is_available", return_value=False):
        out_json = process_document(doc_path, output_path=SAMPLE_DIR / "test_17_chunks.json", enable_semantic=True)

    with open(out_json, "r", encoding="utf-8") as f:
        data = json.load(f)

    assert data["total_chunks"] >= 1
    assert "Section 1: Data Acquisition" in data["chunks"][0]["text"]
    assert "Remote sensing data is downloaded" in data["chunks"][0]["text"]


# =====================================================================
# TEST 18: PDF Processing with Semantic Chunker
# =====================================================================
def test_18_pdf_processing():
    """Verify end-to-end process_document on a native PDF file with semantic chunking."""
    pdf_path = SAMPLE_DIR / "test_18_sem.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(fitz.Point(50, 72), "Chapter 2: Algorithms and Automata\n\nDeterministic finite automata represent computational machines with finite states.")
    doc.save(str(pdf_path))
    doc.close()

    with patch("src.semantic_chunker.OllamaSemanticAdvisor.is_available", return_value=False):
        out_json = process_document(pdf_path, output_path=SAMPLE_DIR / "test_18_chunks.json", enable_ocr=False, enable_semantic=True)

    with open(out_json, "r", encoding="utf-8") as f:
        data = json.load(f)

    assert data["total_chunks"] >= 1
    assert "Algorithms and Automata" in data["chunks"][0]["text"]


# =====================================================================
# TEST 19: OCR Page Processing with Semantic Chunker
# =====================================================================
def test_19_ocr_page_processing():
    """Verify that pages flagged with OCR metadata are chunked properly with ocr_applied=True."""
    chunker = LlamaSemanticChunker(enabled=True)
    pages = [
        DocumentPage(
            page_number=1,
            content="SCANNED_HEADER: Optical character recognition text line.\nSecond scanned line following directly.",
            source="scanned_doc.pdf",
            source_type="scanned_pdf_ocr",
            metadata={"ocr_applied": True, "extraction_method": "OCR"}
        )
    ]

    with patch.object(chunker.advisor, "is_available", return_value=False):
        chunks = chunker.chunk_pages(pages)

    assert len(chunks) >= 1
    assert chunks[0].ocr_applied is True
    assert "SCANNED_HEADER" in chunks[0].text


# =====================================================================
# TEST 20: No HTML Contamination
# =====================================================================
def test_20_no_html_contamination():
    """Verify that chunk.text contains zero HTML tags or UI class markers."""
    chunker = LlamaSemanticChunker(enabled=True)
    text_with_html_leaks = '<div class="chunk-box"><span class="badge">Page 1</span>Authentic prose content here.</div>'
    pages = [DocumentPage(page_number=1, content=text_with_html_leaks, source="doc", source_type="docx")]

    with patch.object(chunker.advisor, "is_available", return_value=False):
        chunks = chunker.chunk_pages(pages)

    for c in chunks:
        assert "<div" not in c.text, f"Found <div in chunk text: {c.text}"
        assert "</div>" not in c.text, f"Found </div> in chunk text: {c.text}"
        assert 'class="chunk-box"' not in c.text
        assert "Authentic prose content here." in c.text
